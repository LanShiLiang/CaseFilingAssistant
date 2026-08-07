from __future__ import annotations

import io

from docx import Document as WordDocument
from PIL import Image, ImageDraw

SYNTHETIC_APPLICANT_ID = "999999199001010016"
SYNTHETIC_RESPONDENT_ID = "999999199202020026"


def legal_basis_docx() -> bytes:
    """生成完全虚构的调解书测试件，999999 行政区划明确表示非真实证件号。"""

    document = WordDocument()
    document.add_paragraph("北京市朝阳区人民法院")
    document.add_paragraph("民事调解书")
    document.add_paragraph("（2026）京0105民初123号")
    document.add_paragraph(
        f"原告：测试原告甲，男，1990年1月1日出生，公民身份号码：{SYNTHETIC_APPLICANT_ID}。"
    )
    document.add_paragraph(
        f"被告：测试被告乙，女，1992年2月2日出生，身份证号：{SYNTHETIC_RESPONDENT_ID}。"
    )
    document.add_paragraph("被执行人应支付人民币 10000.00 元。")
    document.add_paragraph("二〇二六年八月六日")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def synthetic_identity_png(side: str) -> bytes:
    """只画测试占位符，不构造可误用的真实身份证样式或号码。"""

    image = Image.new("RGB", (640, 400), "#edf4f1")
    draw = ImageDraw.Draw(image)
    draw.rectangle((16, 16, 624, 384), outline="#426b61", width=4)
    draw.text((40, 44), f"SYNTHETIC TEST ID - {side}", fill="#173f36")
    draw.text((40, 100), "NOT A REAL IDENTITY DOCUMENT", fill="#b73a33")
    stream = io.BytesIO()
    image.save(stream, format="PNG")
    return stream.getvalue()
