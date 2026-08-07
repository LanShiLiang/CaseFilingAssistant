from __future__ import annotations

import io
from pathlib import Path

from docx import Document as WordDocument

from app.extraction import extract_pages, parse_legal_basis

from .helpers import (
    SYNTHETIC_APPLICANT_ID,
    SYNTHETIC_RESPONDENT_ID,
    legal_basis_docx,
    synthetic_identity_png,
)

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def test_extracts_structured_facts_with_provenance(tmp_path: Path) -> None:
    path = tmp_path / "legal-basis.docx"
    path.write_bytes(legal_basis_docx())

    result = parse_legal_basis(path, DOCX_MIME, "document-test", 20)

    assert result.facts["case_number"] == "（2026）京0105民初123号"
    assert result.facts["document_type"] == "民事调解书"
    assert result.facts["rendering_court"] == "北京市朝阳区人民法院"
    assert result.facts["applicant_name"] == "测试原告甲"
    assert result.facts["applicant_id"] == SYNTHETIC_APPLICANT_ID
    assert result.facts["respondent_name"] == "测试被告乙"
    assert result.facts["respondent_id"] == SYNTHETIC_RESPONDENT_ID
    assert result.facts["judgment_amount"] == "10000.00"
    assert result.sources["case_number"][0]["document_id"] == "document-test"
    assert result.sources["case_number"][0]["confidence"] == 0.92
    assert result.sources["applicant_id"][0]["page"] == 1


def test_rejects_identity_number_with_invalid_checksum(tmp_path: Path) -> None:
    document = WordDocument()
    document.add_paragraph("北京市朝阳区人民法院")
    document.add_paragraph("民事调解书")
    document.add_paragraph("（2026）京0105民初123号")
    document.add_paragraph("原告：测试原告甲，身份证号：999999199001010017。")
    document.add_paragraph("被告：测试被告乙，身份证号：999999199202020027。")
    path = tmp_path / "invalid-identity.docx"
    document.save(path)

    result = parse_legal_basis(path, DOCX_MIME, "invalid-identity-document", 20)

    assert result.facts["applicant_name"] == "测试原告甲"
    assert result.facts["respondent_name"] == "测试被告乙"
    assert "applicant_id" not in result.facts
    assert "respondent_id" not in result.facts


def test_extracts_identity_number_when_pdf_text_inserts_spaces(tmp_path: Path) -> None:
    document = WordDocument()
    document.add_paragraph("北京市朝阳区人民法院")
    document.add_paragraph("民事调解书")
    document.add_paragraph("（2026）京0105民初123号")
    spaced_number = " ".join(SYNTHETIC_APPLICANT_ID)
    document.add_paragraph(f"原告：测试原告甲，公 民 身 份 号 码：{spaced_number}。")
    document.add_paragraph(f"被告：测试被告乙，身份证号：{SYNTHETIC_RESPONDENT_ID}。")
    path = tmp_path / "spaced-identity.docx"
    document.save(path)

    result = parse_legal_basis(path, DOCX_MIME, "spaced-identity-document", 20)

    assert result.facts["applicant_id"] == SYNTHETIC_APPLICANT_ID
    assert "测试原告甲" in result.sources["applicant_id"][0]["snippet"]


def test_image_without_ocr_degrades_to_manual_input(tmp_path: Path, monkeypatch) -> None:
    path = tmp_path / "front.png"
    path.write_bytes(synthetic_identity_png("FRONT"))
    monkeypatch.setattr("app.extraction.image_ocr_available", lambda: False)

    pages, warnings = extract_pages(path, "image/png", 20)

    assert pages[0].method == "ocr"
    assert pages[0].text == ""
    assert warnings == ["image_ocr_unavailable_manual_input_allowed"]


def test_complex_party_signal_is_extracted_and_blocks_silent_scope_drift(
    tmp_path: Path,
) -> None:
    document = WordDocument()
    document.add_paragraph("北京市示例区人民法院")
    document.add_paragraph("民事判决书")
    document.add_paragraph("（2026）示例民初1号")
    document.add_paragraph("申请执行人：测试公司，法定代表人：测试甲")
    stream = io.BytesIO()
    document.save(stream)
    path = tmp_path / "organization.docx"
    path.write_bytes(stream.getvalue())

    result = parse_legal_basis(path, DOCX_MIME, "organization-document", 20)

    codes = {signal["code"] for signal in result.scope_signals}
    assert "unsupported_organization_party" in codes
