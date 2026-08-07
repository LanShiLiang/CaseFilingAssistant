from __future__ import annotations

import io
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document as WordDocument
from PIL import Image, ImageDraw

SYNTHETIC_APPLICANT_ID = "999999199001010016"
SYNTHETIC_RESPONDENT_ID = "999999199202020026"


def write_legal_basis(target: Path) -> None:
    document = WordDocument()
    document.add_paragraph("北京市朝阳区人民法院")
    document.add_paragraph("民事调解书")
    document.add_paragraph("（2026）京0105民初123号")
    document.add_paragraph(
        f"原告：测试原告甲，男，1990年1月1日出生，公民身份号码：{SYNTHETIC_APPLICANT_ID}。"
    )
    document.add_paragraph(
        f"被告：测试被告乙，女，1992年2月2日出生，身份证号：{SYNTHETIC_RESPONDENT_ID}。"
    )
    document.add_paragraph("被执行人应支付人民币 10000.00 元。")
    document.add_paragraph("2026年8月6日")
    document.save(target)


def convert_to_pdf(source: Path, target: Path) -> None:
    office = shutil.which("soffice") or shutil.which("libreoffice")
    if office is None:
        raise RuntimeError("E2E PDF fixture generation requires LibreOffice")
    with tempfile.TemporaryDirectory(prefix="cfa-e2e-pdf-") as directory:
        root = Path(directory)
        profile = root / "lo-profile"
        profile.mkdir()
        completed = subprocess.run(
            [
                office,
                "--headless",
                f"-env:UserInstallation={profile.as_uri()}",
                "--convert-to",
                "pdf",
                "--outdir",
                str(root),
                str(source),
            ],
            capture_output=True,
            text=True,
            timeout=120,
            check=False,
        )
        generated = root / f"{source.stem}.pdf"
        if completed.returncode != 0 or not generated.is_file():
            raise RuntimeError("E2E PDF fixture generation failed")
        target.write_bytes(generated.read_bytes())


def write_identity(target: Path, side: str) -> None:
    image = Image.new("RGB", (640, 400), "#edf4f1")
    draw = ImageDraw.Draw(image)
    draw.rectangle((16, 16, 624, 384), outline="#426b61", width=4)
    draw.text((40, 44), f"SYNTHETIC TEST ID - {side}", fill="#173f36")
    draw.text((40, 100), "NOT A REAL IDENTITY DOCUMENT", fill="#b73a33")
    stream = io.BytesIO()
    image.save(stream, format="PNG")
    target.write_bytes(stream.getvalue())


def main() -> None:
    output = Path(sys.argv[1]).resolve()
    output.mkdir(parents=True, exist_ok=True)
    docx_path = output / "synthetic-legal-basis.docx"
    write_legal_basis(docx_path)
    convert_to_pdf(docx_path, output / "synthetic-legal-basis.pdf")
    write_identity(output / "synthetic-id-front.png", "FRONT")
    write_identity(output / "synthetic-id-back.png", "BACK")


if __name__ == "__main__":
    main()
